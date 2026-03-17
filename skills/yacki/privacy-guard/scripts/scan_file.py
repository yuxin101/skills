import sys
import os
import re

try:
    import pdfplumber
    from docx import Document
    import openpyxl
except ImportError as e:
    requirements_path = os.path.abspath(
        os.path.join(os.path.dirname(__file__), "..", "requirements.txt")
    )
    print(
        "ERROR: Missing Python dependencies. "
        f"Install with: python3 -m pip install -r {requirements_path}"
    )
    sys.exit(2)

# Define sensitive keywords for each category
SENSITIVE_KEYWORDS = {
    "Contract/Agreement": [
        "合同", "协议", "Contract", "Agreement", "NDA", "Non-Disclosure", "保密协议", "采购合同", "销售合同", "劳动合同", "Employment Contract", "Service Agreement", "服务协议", "合作协议", "Partnership Agreement", "Memorandum of Understanding", "MOU", "意向书", "条款", "Terms and Conditions", "Clause", "条款", "盖章", "Seal", "Signature", "签字"
    ],
    "Transaction Data": [
        "交易记录", "支付明细", "Transaction Data", "Amount", "Balance", "余额", "账单", "Invoice", "发票", "收据", "Receipt", "Bank Statement", "银行流水", "转账", "Transfer", "Payment", "支付", "Credit Card", "信用卡", "Debit Card", "借记卡", "Account Number", "账号", "SWIFT", "IBAN", "Routing Number", "Financial Report", "财务报表", "Profit and Loss", "P&L", "损益表", "Balance Sheet", "资产负债表"
    ],
    "Personal Privacy": [
        "身份证", "手机号", "家庭住址", "Personal Privacy", "ID Card", "Phone Number", "Passport", "护照", "Social Security Number", "SSN", "社保号", "Driver's License", "驾照", "Date of Birth", "出生日期", "Email Address", "邮箱地址", "Password", "密码", "Username", "用户名", "Medical Record", "病历", "Health Information", "健康信息", "Biometric", "生物识别", "Fingerprint", "指纹", "Face ID", "人脸识别"
    ],
    "Customer Privacy": [
        "客户信息", "Client Data", "Customer Privacy", "User Profile", "CRM", "Customer Relationship Management", "Mailing List", "邮件列表", "Subscriber", "订阅者", "User ID", "用户ID", "Purchase History", "购买记录", "Behavioral Data", "行为数据", "Personal Identifiable Information", "PII", "GDPR", "Privacy Policy", "隐私政策", "Data Subject", "数据主体", "Consent", "同意书"
    ],
    "Company Confidential": [
        "Confidential", "机密", "Internal Use Only", "内部使用", "Proprietary", "专利", "Trade Secret", "商业秘密", "Business Plan", "商业计划", "Strategy", "战略", "Roadmap", "路线图", "Source Code", "源代码", "Algorithm", "算法", "Architecture", "架构", "Unreleased", "未发布", "Draft", "草案", "Meeting Minutes", "会议纪要", "Board Meeting", "董事会会议", "Salary", "薪资", "Payroll", "工资单", "Employee Evaluation", "员工评估"
    ]
}

def check_text(text):
    """Helper to check text for sensitive keywords."""
    for category, keywords in SENSITIVE_KEYWORDS.items():
        for keyword in keywords:
            if re.search(re.escape(keyword), text, re.IGNORECASE):
                return f"BLOCK: Sensitive keyword '{keyword}' found in content (Category: {category})."
    return None

def scan_file(file_path):
    if not os.path.exists(file_path):
        return f"ERROR: File not found: {file_path}"
    
    filename = os.path.basename(file_path)
    
    # 1. Check filename for sensitive keywords
    for category, keywords in SENSITIVE_KEYWORDS.items():
        for keyword in keywords:
            if keyword.lower() in filename.lower():
                return f"BLOCK: Sensitive keyword '{keyword}' found in filename (Category: {category})."
    
    # 2. Check file content based on extension
    _, ext = os.path.splitext(filename)
    ext = ext.lower()
    
    try:
        # Text-based files
        if ext in {'.txt', '.csv', '.md', '.json', '.xml', '.log', '.py', '.js', '.html'}:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
                result = check_text(content)
                if result: return result
        
        # PDF files
        elif ext == '.pdf':
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    content = page.extract_text() or ""
                    result = check_text(content)
                    if result: return result
        
        # Word documents
        elif ext == '.docx':
            doc = Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            result = check_text("\n".join(full_text))
            if result: return result
        
        # Excel spreadsheets
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    result = check_text(row_text)
                    if result: return result
                    
    except Exception as e:
        return f"ERROR: Could not parse {ext} file: {str(e)}"
    
    return "PASS: No sensitive information detected."

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 scan_file.py <file_path>")
        sys.exit(1)
    
    file_path = sys.argv[1]
    result = scan_file(file_path)
    print(result)
